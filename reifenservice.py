"""Approved tire prices, Decimal calculations and an internal Excel estimate.

No Flask/database/network access and no invoice number or final invoice is
created here. Prices are the gross end prices from
``templates/_homepage_reifenpreise.html``. Keep that view and this versioned
catalogue in sync when an admin approves a future price change.

Controller contract::

    selection = [{"preis_id": "radwechsel_bis17", "menge": 1},
                 {"preis_id": "einlagerung", "menge": 1}]
    calculation = berechne_reifenservice(selection)
    # Save this JSON-compatible snapshot against the order, separately from
    # requested services. An admin selects the services actually performed.
    data = export_reifenservice_xlsx(calculation, kunde={"name": "Max Beispiel"},
                                   auftrag={"id": 123, "kennzeichen": "TEST-1"})

``menge`` counts the displayed unit: Satz = four wheels, Rad/Reifen/Ventil =
one item. Montage already includes wheel removal/refitting and balancing;
overlapping service selections are rejected. ``lose_raeder=True`` applies the
approved 5 EUR per wheel reduction automatically, only to mounting positions.

Gross cents are authoritative. Each position's net amount is derived from its
gross line total, then tax is the remainder. Totals sum those rounded lines,
so every position and the complete estimate reconcile exactly. Do not derive
a two-decimal net unit price and multiply it again when sending to accounting:
use the approved gross unit prices with a gross tax basis instead.
"""

from collections.abc import Mapping
from dataclasses import dataclass, asdict
from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP
from io import BytesIO
import re


PREISKATALOG_VERSION = "reifenpreise-1"
MWST_SATZ = Decimal("19")
_CENT = Decimal("0.01")
_HUNDRED = Decimal("100")
_FAKTOR = Decimal("1.19")
_QUELLE = "Freigegebene Reifenpreisliste"


@dataclass(frozen=True)
class _Preis:
    id: str
    bezeichnung: str
    brutto_cent: int
    einheit: str
    gruppe: str
    hinweis: str = ""


_PREISE = (
    _Preis("radwechsel_bis17", "Radwechsel bis 17 Zoll", 3500, "Satz", "radwechsel", "Vier Kompletträder, ohne Wuchten."),
    _Preis("radwechsel_18_20", "Radwechsel 18–20 Zoll", 4000, "Satz", "radwechsel", "Vier Kompletträder, ohne Wuchten."),
    _Preis("radwechsel_wuchten_bis17", "Radwechsel mit Wuchten bis 17 Zoll", 6000, "Satz", "radwechsel_wuchten", "Vier Räder, einschließlich Standardgewichten."),
    _Preis("radwechsel_wuchten_18_20", "Radwechsel mit Wuchten 18–20 Zoll", 7000, "Satz", "radwechsel_wuchten", "Vier Räder, einschließlich Standardgewichten."),
    _Preis("einlagerung", "Einlagerung bis 20 Zoll", 4500, "Satz", "einlagerung", "Vier Kompletträder, sechs Monate ab Annahme, Kennzeichnung und Sichtcheck; ohne Wäsche."),
    _Preis("montage_13_16", "Reifenmontage 13–16 Zoll", 2250, "Reifen", "montage", "Radaus-/einbau, Ab-/Aufziehen, Wuchten, Standardgewichte und Standard-Gummiventil ohne RDKS."),
    _Preis("montage_17_18", "Reifenmontage 17–18 Zoll", 2500, "Reifen", "montage", "Radaus-/einbau, Ab-/Aufziehen, Wuchten, Standardgewichte und Standard-Gummiventil ohne RDKS."),
    _Preis("montage_19_20", "Reifenmontage 19–20 Zoll", 3000, "Reifen", "montage", "Radaus-/einbau, Ab-/Aufziehen, Wuchten, Standardgewichte und Standard-Gummiventil ohne RDKS."),
    _Preis("saison_radwechsel_bis17", "Einlagerung und Radwechsel bis 17 Zoll", 8000, "Satz", "saison", "Vier Räder, eine Saison / maximal sechs Monate, ohne Wäsche."),
    _Preis("saison_radwechsel_18_20", "Einlagerung und Radwechsel 18–20 Zoll", 8500, "Satz", "saison", "Vier Räder, eine Saison / maximal sechs Monate, ohne Wäsche."),
    _Preis("saison_wuchten_bis17", "Einlagerung und Radwechsel mit Wuchten bis 17 Zoll", 10500, "Satz", "saison", "Vier Räder, eine Saison / maximal sechs Monate, ohne Wäsche."),
    _Preis("saison_wuchten_18_20", "Einlagerung und Radwechsel mit Wuchten 18–20 Zoll", 11500, "Satz", "saison", "Vier Räder, eine Saison / maximal sechs Monate, ohne Wäsche."),
    _Preis("raederwaesche", "Räderwäsche", 1500, "Satz", "waesche", "Vier Räder, normale Verschmutzung."),
    _Preis("wuchten_los_bis17", "Wuchten loser Räder bis 17 Zoll", 625, "Rad", "wuchten", "Einschließlich Gewichten."),
    _Preis("wuchten_los_18_20", "Wuchten loser Räder 18–20 Zoll", 750, "Rad", "wuchten", "Einschließlich Gewichten."),
    _Preis("runflat", "Runflat / Niederquerschnitt bei Reifenmontage", 500, "Reifen", "montage_zusatz", "Einmal je Reifen; nicht beim reinen Komplettradwechsel."),
    _Preis("entsorgung", "Altreifenentsorgung ohne Felge", 400, "Reifen", "entsorgung"),
    _Preis("metallventil", "Metallventil statt Standard-Gummiventil", 450, "Ventil", "montage_zusatz", "Ohne RDKS. Ersetzt das enthaltene Gummiventil, kein RDKS-Servicekit."),
)
_PREIS_NACH_ID = {preis.id: preis for preis in _PREISE}
_LOSE_RAEDER = _Preis("lose_raeder", "Nachlass bei angelieferten losen Rädern", -500, "Rad", "nachlass", "5 EUR weniger je montiertem Rad.")


def reifen_preiskatalog():
    """JSON-safe copies of the selectable gross prices, never shared state."""
    return [asdict(preis) | {"mwst_satz": "19.00", "katalog_version": PREISKATALOG_VERSION}
            for preis in _PREISE]


def _menge(value, preis):
    if isinstance(value, bool) or not isinstance(value, (str, int)) or not re.fullmatch(r"[1-4]", str(value)):
        raise ValueError("Bitte eine ganze Menge von 1 bis 4 angeben.")
    menge = int(value)
    if preis.einheit == "Satz" and menge != 1:
        raise ValueError("Ein Satz umfasst vier Räder. Pro Reifenauftrag ist ein Satz auswählbar.")
    return menge


def _position(preis, menge):
    brutto = (Decimal(preis.brutto_cent) * menge / _HUNDRED).quantize(_CENT)
    netto = (brutto / _FAKTOR).quantize(_CENT, rounding=ROUND_HALF_UP)
    return {
        "preis_id": preis.id, "bezeichnung": preis.bezeichnung,
        "menge": menge, "einheit": preis.einheit, "hinweis": preis.hinweis,
        "einzel_brutto_cent": preis.brutto_cent,
        "brutto_cent": int(brutto * _HUNDRED),
        "netto_cent": int(netto * _HUNDRED),
        "mwst_cent": int((brutto - netto) * _HUNDRED),
    }


def berechne_reifenservice(auswahl, *, lose_raeder=False):
    """Validate actual-service selections and return an exact JSON snapshot.

    At most one four-wheel vehicle/set per calculation. Mixed mounting or loose
    balancing sizes can use several lines, up to four tires/wheels in total.
    Unknown prices/fields, repeated IDs, overlapping packages and invalid
    quantities raise ValueError. Zero quantities should be omitted by the UI.
    """
    if not isinstance(lose_raeder, bool):
        raise ValueError("Lose Räder bitte ausdrücklich mit Ja oder Nein angeben.")
    if not isinstance(auswahl, (list, tuple)) or not auswahl or len(auswahl) > len(_PREISE):
        raise ValueError("Bitte mindestens eine ausgeführte Reifenleistung auswählen.")
    gewaehlt = []
    gesehen = set()
    for item in auswahl:
        if not isinstance(item, Mapping) or set(item) - {"preis_id", "menge"}:
            raise ValueError("Eine Leistungsposition muss Preis-ID und Menge enthalten.")
        preis_id = item.get("preis_id")
        if not isinstance(preis_id, str) or preis_id not in _PREIS_NACH_ID:
            raise ValueError("Unbekannte Reifenleistung. Bitte die aktuelle Preisliste verwenden.")
        if preis_id in gesehen:
            raise ValueError("Jede Reifenleistung darf nur einmal ausgewählt werden.")
        gesehen.add(preis_id)
        preis = _PREIS_NACH_ID[preis_id]
        gewaehlt.append((preis, _menge(item.get("menge"), preis)))
    gruppen = {preis.gruppe for preis, _ in gewaehlt}
    haupt = [(preis, menge) for preis, menge in gewaehlt if preis.gruppe in {"radwechsel", "radwechsel_wuchten", "saison"}]
    if len(haupt) > 1:
        raise ValueError("Bitte nur einen Radwechsel- oder Saisonpaketpreis auswählen.")
    if (haupt and gruppen & {"montage", "wuchten"}) or ("montage" in gruppen and "wuchten" in gruppen):
        raise ValueError("Wuchten und Radwechsel sind im gewählten Paket bereits enthalten oder separat anders bepreist.")
    if "saison" in gruppen and "einlagerung" in gruppen:
        raise ValueError("Die Einlagerung ist im Saisonpaket bereits enthalten.")
    montage_menge = sum(menge for preis, menge in gewaehlt if preis.gruppe == "montage")
    wuchten_menge = sum(menge for preis, menge in gewaehlt if preis.gruppe == "wuchten")
    if montage_menge > 4 or wuchten_menge > 4:
        raise ValueError("Ein Reifenauftrag umfasst höchstens vier Räder beziehungsweise Reifen.")
    for preis, menge in gewaehlt:
        if preis.gruppe == "montage_zusatz" and (not montage_menge or menge > montage_menge):
            raise ValueError("Montage-Zuschläge nur bis zur Anzahl der tatsächlich montierten Reifen auswählen.")
    if lose_raeder and not montage_menge:
        raise ValueError("Der Nachlass für lose Räder gilt nur bei Reifenmontage.")
    positionen = [_position(preis, menge) for preis, menge in gewaehlt]
    if lose_raeder:
        positionen.append(_position(_LOSE_RAEDER, montage_menge))
    return {
        "schema_version": 1, "katalog_version": PREISKATALOG_VERSION,
        "quelle": _QUELLE, "status": "entwurf", "steuerbasis": "brutto",
        "mwst_satz": "19.00", "lose_raeder": lose_raeder,
        "positionen": positionen,
        "gesamt": {key: sum(position[key] for position in positionen)
                   for key in ("brutto_cent", "netto_cent", "mwst_cent")},
    }


def _gepruefter_snapshot(kalkulation):
    if not isinstance(kalkulation, Mapping):
        raise ValueError("Eine gespeicherte Reifen-Kalkulation fehlt.")
    positionen = kalkulation.get("positionen")
    if not isinstance(positionen, list) or any(not isinstance(p, Mapping) for p in positionen):
        raise ValueError("Die gespeicherten Leistungspositionen sind ungültig.")
    auswahl = [{"preis_id": p.get("preis_id"), "menge": p.get("menge")}
               for p in positionen if p.get("preis_id") != "lose_raeder"]
    aktuell = berechne_reifenservice(auswahl, lose_raeder=kalkulation.get("lose_raeder", False))
    if aktuell != dict(kalkulation):
        raise ValueError("Die Kalkulation stimmt nicht mit den gespeicherten Preispositionen überein. Bitte neu prüfen.")
    return aktuell


def _text(value):
    # Explicit string cells prevent Excel formula injection in customer input.
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", str(value or ""))
    return " ".join(text.split())[:300]


def export_reifenservice_xlsx(kalkulation, *, kunde=None, auftrag=None):
    """Return printable XLSX bytes of a validated calculation, never an invoice.

    Customer keys: name, strasse, plz, ort, telefon, email. Order keys: id,
    kennzeichen, fahrzeug, leistungsdatum. This function does not assign a
    document number, change order state, contact accounting or send anything.
    ``openpyxl`` is deliberately imported only for this optional export.
    """
    kalkulation = _gepruefter_snapshot(kalkulation)
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    except ImportError as exc:
        raise RuntimeError("Der Excel-Export benötigt die installierte openpyxl-Laufzeit.") from exc
    kunde = kunde if isinstance(kunde, Mapping) else {}
    auftrag = auftrag if isinstance(auftrag, Mapping) else {}
    workbook = Workbook()
    workbook.properties.title = "Reifenservice-Kalkulation (Entwurf)"
    workbook.properties.subject = "Interne Leistungs- und Preiskalkulation, keine Rechnung"
    sheet = workbook.active
    sheet.title = "Reifenservice"
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.sheet_properties.tabColor = "26343F"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_options.horizontalCentered = True
    sheet.oddFooter.center.text = "Kalkulation / Entwurf – keine Rechnung"

    def text(zelle, wert):
        cell = sheet[zelle]
        cell.value = _text(wert)
        cell.data_type = "s"

    text("A2", "Reifenservice-Kalkulation")
    sheet["A2"].font = Font(name="Arial", size=14, bold=True)
    text("A3", "ENTWURF – keine Rechnung")
    sheet["A3"].font = Font(name="Arial", size=10, bold=True, color="B42318")
    text("A5", "Kunde")
    text("B5", kunde.get("name"))
    text("A6", "Anschrift")
    text("B6", " ".join(filter(None, (_text(kunde.get("strasse")), _text(kunde.get("plz")), _text(kunde.get("ort"))))))
    text("A7", "Kontakt")
    text("B7", " / ".join(filter(None, (_text(kunde.get("telefon")), _text(kunde.get("email"))))))
    text("F5", "Auftrag")
    text("G5", auftrag.get("id"))
    text("F6", "Fahrzeug")
    text("G6", " / ".join(filter(None, (_text(auftrag.get("fahrzeug")), _text(auftrag.get("kennzeichen"))))))
    text("F7", "Leistungsdatum")
    leistungsdatum = auftrag.get("leistungsdatum")
    if isinstance(leistungsdatum, (date, datetime)):
        sheet["G7"] = leistungsdatum
    elif isinstance(leistungsdatum, str) and leistungsdatum:
        try:
            sheet["G7"] = datetime.strptime(leistungsdatum, "%d.%m.%Y").date() if "." in leistungsdatum else date.fromisoformat(leistungsdatum)
        except ValueError as exc:
            raise ValueError("Bitte ein gültiges Leistungsdatum angeben.") from exc
    sheet["G7"].number_format = "dd.mm.yyyy"
    sheet["G6"].alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[6].height = 36
    text("A9", "Endpreise inklusive 19 % MwSt. Ein Satz umfasst vier Räder.")
    text("A10", f"Preisgrundlage: {_QUELLE}, {PREISKATALOG_VERSION}.")
    for cell in ("A9", "A10"):
        sheet[cell].font = Font(name="Arial", size=10, italic=True, color="667085")

    headers = ("Leistung", "Menge", "Einheit", "Einzel brutto (€)", "Netto (€)", "MwSt. (€)", "Brutto (€)")
    for col, label in enumerate(headers, start=1):
        cell = sheet.cell(12, col, label)
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="26343F")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row, position in enumerate(kalkulation["positionen"], start=13):
        text(f"A{row}", position["bezeichnung"])
        if position["hinweis"]:
            sheet.cell(row, 1).value += "\n" + position["hinweis"]
        sheet.cell(row, 2, position["menge"])
        text(f"C{row}", position["einheit"])
        for col, key in ((4, "einzel_brutto_cent"), (5, "netto_cent"), (6, "mwst_cent"), (7, "brutto_cent")):
            sheet.cell(row, col, Decimal(position[key]) / _HUNDRED)
        sheet.cell(row, 1).alignment = Alignment(wrap_text=True, vertical="center")
        sheet.row_dimensions[row].height = 16 * (sum(max(1, (len(part) + 54) // 55) for part in sheet.cell(row, 1).value.splitlines())) + 6
    last = 12 + len(kalkulation["positionen"])
    total = last + 2
    text(f"A{total}", "Gesamt")
    for col, key in ((5, "netto_cent"), (6, "mwst_cent"), (7, "brutto_cent")):
        sheet.cell(total, col, Decimal(kalkulation["gesamt"][key]) / _HUNDRED)
    for cell in sheet[total][:7]:
        cell.font = Font(name="Arial", size=10, bold=True)
        cell.border = Border(top=Side(style="thin", color="667085"))
    text(f"A{total + 2}", "Die endgültige Rechnung wird im Rechnungssystem erstellt.")
    sheet.print_area = f"A1:G{total + 2}"
    sheet.print_title_rows = "12:12"
    widths = {"A": 58, "B": 9, "C": 12, "D": 20, "E": 17, "F": 17, "G": 18}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    for row in sheet.iter_rows(min_row=5, max_row=total + 2, max_col=7):
        for cell in row:
            if cell.row not in (9, 10, 12, total):
                cell.font = Font(name="Arial", size=10)
            if cell.row >= 13 and cell.column in (2, 3):
                cell.alignment = Alignment(horizontal="right" if cell.column == 2 else "left", vertical="center")
            if cell.row >= 13 and cell.column >= 4 and isinstance(cell.value, (int, float, Decimal)):
                cell.number_format = '#,##0.00'
                cell.alignment = Alignment(horizontal="right", vertical="center")
    sheet.row_dimensions[12].height = 30
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
