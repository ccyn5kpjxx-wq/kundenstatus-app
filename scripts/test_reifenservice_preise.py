"""Isolated approved-price, double-charge and XLSX regression checks.

Run ``python scripts/test_reifenservice_preise.py``. Uses synthetic customer
data, in-memory workbooks and no portal import, database or network access.
"""

from copy import deepcopy
from datetime import datetime
from io import BytesIO
from pathlib import Path
import json
import re
import sys
import unittest
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from reifenservice import (PREISKATALOG_VERSION, berechne_reifenservice,
                          export_reifenservice_xlsx, reifen_preiskatalog)


def calc(*rows, **options):
    return berechne_reifenservice([{"preis_id": key, "menge": count} for key, count in rows], **options)


class ReifenservicePreiseTest(unittest.TestCase):
    def test_catalogue_only_contains_published_prices(self):
        template = (Path(__file__).resolve().parents[1] / "templates/_homepage_reifenpreise.html").read_text(encoding="utf-8")
        published_cents = {int(value) for value in re.findall(r"price\((\d+)\)", template)}
        published_cents.update(int(value) * 100 for value in re.findall(r"(?:>|strong>)(\d+) €", template))
        catalogue = reifen_preiskatalog()
        self.assertEqual(len(catalogue), 18)
        values = {p["brutto_cent"] for p in catalogue}
        values.update(4 * p["brutto_cent"] for p in catalogue if p["gruppe"] == "montage")
        self.assertEqual(values, published_cents)
        self.assertIn("Alle Preise inklusive 19 % MwSt.", template)
        self.assertTrue(all(p["mwst_satz"] == "19.00" for p in catalogue))
        self.assertTrue(all(p["katalog_version"] == PREISKATALOG_VERSION for p in catalogue))

    def test_catalogue_is_a_copy(self):
        first = reifen_preiskatalog()
        first[0]["brutto_cent"] = 1
        self.assertEqual(reifen_preiskatalog()[0]["brutto_cent"], 3500)

    def test_radwechsel_and_storage_are_exact_80_euro(self):
        result = calc(("radwechsel_bis17", 1), ("einlagerung", 1))
        self.assertEqual(result["gesamt"], {"brutto_cent": 8000, "netto_cent": 6723, "mwst_cent": 1277})
        self.assertEqual(result["status"], "entwurf")
        self.assertEqual(result["steuerbasis"], "brutto")
        self.assertEqual(result, json.loads(json.dumps(result)))

    def test_mounting_rounds_the_line_not_net_unit_price(self):
        result = calc(("montage_13_16", 4))
        self.assertEqual(result["gesamt"], {"brutto_cent": 9000, "netto_cent": 7563, "mwst_cent": 1437})
        self.assertEqual(result["positionen"][0]["einzel_brutto_cent"], 2250)

    def test_mixed_sizes_and_loose_wheel_discount(self):
        result = calc(("montage_13_16", 2), ("montage_19_20", 2), ("runflat", 2),
                      ("metallventil", 4), ("entsorgung", 4), lose_raeder=True)
        self.assertEqual(result["gesamt"]["brutto_cent"], 12900)
        self.assertEqual([p["einzel_brutto_cent"] for p in result["positionen"][:2]], [1750,2500])
        self.assertIn('Nachlass',result["positionen"][0]["hinweis"])
        for position in result["positionen"]:
            self.assertEqual(position["netto_cent"] + position["mwst_cent"], position["brutto_cent"])
        self.assertEqual(result["gesamt"]["netto_cent"] + result["gesamt"]["mwst_cent"], 12900)

    def test_loose_wheel_discount_uses_discounted_gross_vat_basis(self):
        result = calc(("montage_19_20",3),lose_raeder=True)
        self.assertEqual(result["gesamt"],{"brutto_cent":7500,"netto_cent":6303,"mwst_cent":1197})

    def test_all_season_package_prices(self):
        for key, gross in (("saison_radwechsel_bis17", 8000), ("saison_radwechsel_18_20", 8500),
                           ("saison_wuchten_bis17", 10500), ("saison_wuchten_18_20", 11500)):
            with self.subTest(key=key):
                self.assertEqual(calc((key, 1))["gesamt"]["brutto_cent"], gross)

    def test_extras_have_the_published_unit_basis(self):
        result = calc(("wuchten_los_bis17", 4), ("raederwaesche", 1), ("entsorgung", 4))
        self.assertEqual(result["gesamt"]["brutto_cent"], 5600)
        self.assertEqual([p["einheit"] for p in result["positionen"]], ["Rad", "Satz", "Reifen"])

    def test_invalid_quantities_ids_and_payloads_rejected(self):
        for value in (None, 0, -1, 5, True, False, "1.5", "1,5", " 1 ", 1.0, "NaN", "1e0"):
            with self.subTest(value=value), self.assertRaises(ValueError):
                calc(("montage_13_16", value))
        for payload in (None, {}, [], ["radwechsel_bis17"], [{"preis_id": "unbekannt", "menge": 1}],
                        [{"preis_id": [], "menge": 1}], [{"preis_id": "radwechsel_bis17", "menge": 1, "brutto_cent": 1}]):
            with self.subTest(payload=payload), self.assertRaises(ValueError):
                berechne_reifenservice(payload)
        with self.assertRaises(ValueError):
            calc(("radwechsel_bis17", 2))
        with self.assertRaises(ValueError):
            calc(("montage_13_16", 1), lose_raeder="1")

    def test_duplicate_and_overlapping_services_rejected(self):
        conflicts = [
            (("radwechsel_bis17", 1), ("radwechsel_bis17", 1)),
            (("radwechsel_bis17", 1), ("radwechsel_18_20", 1)),
            (("radwechsel_bis17", 1), ("radwechsel_wuchten_bis17", 1)),
            (("radwechsel_bis17", 1), ("wuchten_los_bis17", 4)),
            (("montage_13_16", 4), ("radwechsel_bis17", 1)),
            (("montage_17_18", 4), ("wuchten_los_bis17", 4)),
            (("saison_wuchten_bis17", 1), ("einlagerung", 1)),
            (("saison_radwechsel_bis17", 1), ("montage_13_16", 4)),
            (("saison_radwechsel_bis17", 1), ("wuchten_los_bis17", 4)),
            (("montage_13_16", 3), ("montage_17_18", 2)),
            (("wuchten_los_bis17", 3), ("wuchten_los_18_20", 2)),
            (("radwechsel_bis17", 1), ("runflat", 4)),
            (("montage_17_18", 2), ("metallventil", 3)),
        ]
        for rows in conflicts:
            with self.subTest(rows=rows), self.assertRaises(ValueError):
                calc(*rows)
        with self.assertRaises(ValueError):
            calc(("radwechsel_bis17", 1), lose_raeder=True)

    def test_export_rejects_tampered_price_and_version(self):
        original = calc(("montage_13_16", 4), lose_raeder=True)
        for mutate in (lambda p: p["gesamt"].update(brutto_cent=1),
                       lambda p: p["positionen"][0].update(einzel_brutto_cent=1),
                       lambda p: p.update(katalog_version="andere-preise"),
                       lambda p: p.update(status="rechnung"),
                       lambda p: p["positionen"].pop()):
            modified = deepcopy(original)
            mutate(modified)
            with self.subTest(modified=modified), self.assertRaises(ValueError):
                export_reifenservice_xlsx(modified)

    def test_export_has_typed_values_and_is_clearly_not_invoice(self):
        from openpyxl import load_workbook
        result = calc(("radwechsel_bis17", 1), ("einlagerung", 1))
        data = export_reifenservice_xlsx(result,
              kunde={"name": "=HYPERLINK(\"https://example.test\",\"x\")", "strasse": "Teststraße 1", "plz": "01234", "ort": "Testort"},
              auftrag={"id": 123, "fahrzeug": "Testfahrzeug", "kennzeichen": "TEST-123", "leistungsdatum": "20.10.2026"})
        self.assertTrue(data.startswith(b"PK"))
        workbook = load_workbook(BytesIO(data), data_only=False)
        sheet = workbook.active
        self.assertEqual(workbook.sheetnames, ["Reifenservice"])
        self.assertEqual(sheet["A3"].value, "ENTWURF – keine Rechnung")
        self.assertEqual(sheet["B5"].data_type, "s")
        self.assertTrue(sheet["B5"].value.startswith("=HYPERLINK"))
        self.assertEqual(sheet["G16"].value, 80)
        self.assertEqual(sheet["E16"].value, 67.23)
        self.assertEqual(sheet["F16"].value, 12.77)
        self.assertEqual(sheet["B13"].value, 1)
        self.assertEqual(sheet["G13"].value, 35)
        self.assertTrue(sheet["G13"].data_type == "n")
        self.assertFalse(sheet.sheet_view.showGridLines)
        self.assertIn("A1:G18", str(sheet.print_area).replace("$", ""))
        self.assertEqual(sheet.page_setup.orientation, "landscape")
        self.assertEqual(sheet["G7"].value, datetime(2026, 10, 20))
        self.assertIn("sechs Monate", sheet["A14"].value)
        self.assertFalse(any(cell.data_type == "f" for row in sheet for cell in row))
        self.assertFalse(any("Rechnungsnummer" in str(cell.value) for row in sheet for cell in row))
        workbook.close()

    def test_export_works_with_no_optional_customer_values(self):
        from openpyxl import load_workbook
        data = export_reifenservice_xlsx(calc(("entsorgung", 1)))
        workbook = load_workbook(BytesIO(data))
        self.assertEqual(workbook.active["G15"].value, 4)
        workbook.close()

    def test_export_rejects_invalid_service_date(self):
        with self.assertRaises(ValueError):
            export_reifenservice_xlsx(calc(("einlagerung", 1)), auftrag={"leistungsdatum": "31.02.2026"})

    def test_export_normalizes_xml_control_characters(self):
        from openpyxl import load_workbook
        workbook = load_workbook(BytesIO(export_reifenservice_xlsx(calc(("einlagerung", 1)), kunde={"name": "Max\x00Beispiel\x08"})))
        self.assertEqual(workbook.active["B5"].value, "MaxBeispiel")
        workbook.close()


if __name__ == "__main__":
    unittest.main(verbosity=2)
